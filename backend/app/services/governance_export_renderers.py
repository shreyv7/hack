"""Format-specific renderers for governance dashboard exports."""

from __future__ import annotations

import csv
import io
import json
from typing import Any

from fastapi import HTTPException, status
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.governance_dashboard import ExportFormat


def _flatten_metrics(data: dict[str, Any], prefix: str = "") -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    for key, value in data.items():
        path = f"{prefix}.{key}" if prefix else key
        if isinstance(value, dict):
            rows.extend(_flatten_metrics(value, path))
        elif isinstance(value, list):
            rows.append((path, json.dumps(value, default=str)))
        else:
            rows.append((path, str(value)))
    return rows


def render_json(payload: dict[str, Any]) -> bytes:
    return json.dumps(payload, default=str, indent=2).encode()


def render_csv(dashboard_name: str, filters: dict[str, Any], widget_payloads: list[dict[str, Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["dashboard", dashboard_name])
    writer.writerow(["time_range", filters.get("time_range", "90d")])
    writer.writerow(["business_unit_id", filters.get("business_unit_id", "")])
    writer.writerow([])
    writer.writerow(["widget_title", "data_source", "metric_key", "metric_value"])
    for item in widget_payloads:
        widget = item.get("widget") or {}
        metrics = item.get("metrics") or {}
        data = metrics.get("data") if isinstance(metrics, dict) else {}
        title = widget.get("title", "")
        source = widget.get("data_source", "")
        if isinstance(source, dict):
            source = source.get("value", "")
        for key, value in _flatten_metrics(data if isinstance(data, dict) else {}):
            writer.writerow([title, source, key, value])
    return buffer.getvalue().encode()


def render_pdf(dashboard_name: str, description: str | None, widget_payloads: list[dict[str, Any]]) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=36)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SectionTitle", parent=styles["Heading2"], spaceAfter=8))
    story: list[Any] = []

    story.append(Paragraph(f"Compliance Dashboard: {dashboard_name}", styles["Title"]))
    if description:
        story.append(Paragraph(description, styles["Normal"]))
    story.append(Spacer(1, 12))

    summary_rows = [["Widget", "Data Source", "Key Metrics"]]
    for item in widget_payloads:
        widget = item.get("widget") or {}
        metrics = item.get("metrics") or {}
        data = metrics.get("data") if isinstance(metrics, dict) else {}
        flat = _flatten_metrics(data if isinstance(data, dict) else {})
        preview = "; ".join(f"{k}: {v}" for k, v in flat[:3]) or "—"
        source = widget.get("data_source", "")
        if isinstance(source, dict):
            source = source.get("value", "")
        summary_rows.append([widget.get("title", ""), str(source), preview])

    table = Table(summary_rows, colWidths=[2.2 * inch, 1.5 * inch, 3.3 * inch])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a5f")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 18))

    for item in widget_payloads:
        widget = item.get("widget") or {}
        metrics = item.get("metrics") or {}
        data = metrics.get("data") if isinstance(metrics, dict) else {}
        story.append(Paragraph(widget.get("title", "Widget"), styles["SectionTitle"]))
        flat = _flatten_metrics(data if isinstance(data, dict) else {})
        if not flat:
            story.append(Paragraph("No metrics available.", styles["Normal"]))
        else:
            detail_rows = [["Metric", "Value"]] + [[k, v[:200]] for k, v in flat[:25]]
            detail_table = Table(detail_rows, colWidths=[2.5 * inch, 4.5 * inch])
            detail_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f7")),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                    ]
                )
            )
            story.append(detail_table)
        story.append(Spacer(1, 12))

    doc.build(story)
    buffer.seek(0)
    return buffer.read()


def render_xlsx(dashboard_name: str, widget_payloads: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Dashboard", dashboard_name])
    summary.append(["Widget", "Data Source"])
    for item in widget_payloads:
        widget = item.get("widget") or {}
        source = widget.get("data_source", "")
        if isinstance(source, dict):
            source = source.get("value", "")
        summary.append([widget.get("title", ""), str(source)])

    for idx, item in enumerate(widget_payloads, start=1):
        widget = item.get("widget") or {}
        metrics = item.get("metrics") or {}
        data = metrics.get("data") if isinstance(metrics, dict) else {}
        title = (widget.get("title") or f"Widget {idx}")[:28]
        ws = wb.create_sheet(title=title)
        ws.append(["Metric", "Value"])
        for key, value in _flatten_metrics(data if isinstance(data, dict) else {}):
            ws.append([key, value])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def render_pptx(dashboard_name: str, widget_payloads: list[dict[str, Any]]) -> bytes:
    try:
        from pptx import Presentation
        from pptx.util import Pt
    except ImportError as exc:
        raise HTTPException(
            status.HTTP_501_NOT_IMPLEMENTED,
            detail="PPTX export requires python-pptx on the server",
        ) from exc

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = dashboard_name
    title_slide.placeholders[1].text = "DPDPA Shield — Governance Dashboard Export"

    for item in widget_payloads:
        widget = item.get("widget") or {}
        metrics = item.get("metrics") or {}
        data = metrics.get("data") if isinstance(metrics, dict) else {}
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = widget.get("title", "Widget")
        body = slide.placeholders[1].text_frame
        body.clear()
        flat = _flatten_metrics(data if isinstance(data, dict) else {})[:12]
        if not flat:
            p = body.paragraphs[0]
            p.text = "No metrics available."
        else:
            for i, (key, value) in enumerate(flat):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = f"{key}: {value}"
                p.font.size = Pt(14)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


CONTENT_TYPES: dict[ExportFormat, str] = {
    ExportFormat.json: "application/json",
    ExportFormat.csv: "text/csv",
    ExportFormat.pdf: "application/pdf",
    ExportFormat.xlsx: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ExportFormat.pptx: "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def render_export(
    fmt: ExportFormat,
    *,
    slug: str,
    dashboard_name: str,
    description: str | None,
    filters: dict[str, Any],
    payload: dict[str, Any],
    widget_payloads: list[dict[str, Any]],
) -> tuple[bytes, str, str]:
    if fmt == ExportFormat.png:
        raise HTTPException(status.HTTP_501_NOT_IMPLEMENTED, detail="PNG export is not yet supported")
    if fmt == ExportFormat.json:
        raw = render_json(payload)
    elif fmt == ExportFormat.csv:
        raw = render_csv(dashboard_name, filters, widget_payloads)
    elif fmt == ExportFormat.pdf:
        raw = render_pdf(dashboard_name, description, widget_payloads)
    elif fmt == ExportFormat.xlsx:
        raw = render_xlsx(dashboard_name, widget_payloads)
    elif fmt == ExportFormat.pptx:
        raw = render_pptx(dashboard_name, widget_payloads)
    else:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=f"Unsupported export format: {fmt}")

    content_type = CONTENT_TYPES.get(fmt, "application/octet-stream")
    filename = f"{slug}.{fmt.value}"
    return raw, filename, content_type
